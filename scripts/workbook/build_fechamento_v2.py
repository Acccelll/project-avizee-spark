#!/usr/bin/env python3
"""Gera public/workbook_fechamento_v2.xlsx, o modelo do Workbook de Fechamento.

Uso:
    python3 scripts/workbook/build_fechamento_v2.py [saida.xlsx]

O modelo traz 12 abas com a identidade da AviZee. As 11 abas de análise só
têm fórmulas que leem a aba Dados; o sistema preenche a aba Dados na geração
(src/lib/workbook/fechamento/v2/) e grava o valor calculado de cada fórmula.
As fórmulas usam só funções que o avaliador do sistema conhece
(src/lib/workbook/fechamento/formulas.ts): IF, IFERROR, SUM, AVERAGE, SUMIFS,
AND, OR, ABS, ROUND, NA, ISNUMBER, N, EOMONTH, VLOOKUP e LOOKUP.

Layout da aba Dados (precisa bater com src/lib/workbook/fechamento/v2/layout.ts):
  Mensal, cabeçalho na linha 4, linhas 5..41 = dez/(ano-3) .. dez/(ano)
    jan/(ano-2) = 6, jan/(ano-1) = 18, jan/(ano) = 30
    A mês (data) | B rótulo ("set/26") | C 1 se mês fechado | D recebido | E pago
    F faturamento | G caixa final (total) | H caixa calculado pelo ERP
    I ajuste ao extrato | J bloqueado | K estoque produtos | L estoque materiais
    M..Q a receber: a vencer até 30, 31-90, >90, vencido até 30, >30
    R..V a pagar: mesmas faixas | W LinkedIn | X Instagram
    Y..AB retirada sócio 1..4 | AC..AF resultado do ano atribuído sócio 1..4
    AG rendimentos de aplicações (já incluídos no recebido)
  Parâmetros AH (rótulo) / AI (valor), linhas 5..28 (ver PARAM abaixo)
  Listas (linhas 5..44): AK cliente, AL 1ª compra, AN UF, AO estado,
    AQ categoria de despesa, AS fornecedor do mês
  Detalhe, cabeçalho na linha 46, a partir da 47:
    A mês | B rótulo | C tipo (Cliente, Estado, Despesa, Fornecedor)
    D item | E chave | F valor | G notas
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.data_source import AxDataSource, NumFmt
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import SeriesLabel, StrRef
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.image import Image
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties

RAIZ = Path(__file__).resolve().parents[2]
LOGO = RAIZ / 'public/images/logoavizee.png'

# ── Marca ────────────────────────────────────────────────────────────────────
VINHO, VINHO2, VINHO3 = '690500', 'A0473D', 'DCBCB3'
TERRA, TERRA2, TERRA3 = 'B2592C', 'D08A63', 'ECC6AC'
BEGE, BEGE_CLARO, CREME = 'DDD2C4', 'E9E0D4', 'FFFAED'
TINTA, APAGADO, REGRA = '1B1411', '8A7E73', 'E4DCD2'
BOM, RUIM, BRANCO = '4F7A46', 'B3261E', 'FFFFFF'
FONTE = 'Montserrat'

BRL = '#,##0;[Red](#,##0);"–"'
BRL2 = '#,##0.00;[Red]-#,##0.00;"–"'
PCT = '0.0%;[Red]-0.0%;"–"'
INT = '#,##0;[Red]-#,##0;"–"'
MESES_TXT = '0.0;[Red]-0.0;"–"'
DATA_MES = 'mmm/yy'
DATA = 'dd/mm/yyyy'

MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']

# ── Aba Dados ────────────────────────────────────────────────────────────────
D_CAB, D_INI, D_FIM = 4, 5, 41
JAN = {-2: 6, -1: 18, 0: 30}
SLOTS_LISTA = 40
DET_CAB, DET_INI, DET_FIM = 46, 47, 4000

COLS = [
    ('A', 'Mês', DATA_MES), ('B', 'Rótulo', None), ('C', 'Fechado', '0'),
    ('D', 'Recebido', BRL2), ('E', 'Pago', BRL2), ('F', 'Faturamento', BRL2),
    ('G', 'Caixa final', BRL2), ('H', 'Caixa calculado (ERP)', BRL2), ('I', 'Ajuste ao extrato', BRL2),
    ('J', 'Bloqueado', BRL2), ('K', 'Estoque produtos', BRL2), ('L', 'Estoque materiais', BRL2),
    ('M', 'Receber a vencer 0-30', BRL2), ('N', 'Receber a vencer 31-90', BRL2), ('O', 'Receber a vencer >90', BRL2),
    ('P', 'Receber vencido 0-30', BRL2), ('Q', 'Receber vencido >30', BRL2),
    ('R', 'Pagar a vencer 0-30', BRL2), ('S', 'Pagar a vencer 31-90', BRL2), ('T', 'Pagar a vencer >90', BRL2),
    ('U', 'Pagar vencido 0-30', BRL2), ('V', 'Pagar vencido >30', BRL2),
    ('W', 'Seguidores LinkedIn', INT), ('X', 'Seguidores Instagram', INT),
    ('Y', 'Retirada sócio 1', BRL2), ('Z', 'Retirada sócio 2', BRL2), ('AA', 'Retirada sócio 3', BRL2), ('AB', 'Retirada sócio 4', BRL2),
    ('AC', 'Resultado atribuído sócio 1', BRL2), ('AD', 'Resultado atribuído sócio 2', BRL2),
    ('AE', 'Resultado atribuído sócio 3', BRL2), ('AF', 'Resultado atribuído sócio 4', BRL2),
    ('AG', 'Rendimentos de aplicações', BRL2),
]
IDX = {c: i + 1 for i, (c, _, _) in enumerate(COLS)}  # coluna -> índice no VLOOKUP a partir de A

PARAM = [
    ('Competência', DATA_MES), ('Mês', '0'), ('Ano', '0'), ('Período', None), ('Data-base', DATA),
    ('Gerado em', DATA), ('Meta de crescimento do ano', PCT), ('Meta de crescimento do ano anterior', PCT),
    ('Limite de faturamento (ano-2)', BRL2), ('Limite de faturamento (ano-1)', BRL2), ('Limite de faturamento (ano)', BRL2),
    ('Linhas do extrato pendentes', '0'), ('Último extrato importado', DATA), ('Notas de saída no mês', '0'),
    ('Sócio 1', None), ('Sócio 2', None), ('Sócio 3', None), ('Sócio 4', None),
    ('Participação sócio 1', PCT), ('Participação sócio 2', PCT), ('Participação sócio 3', PCT), ('Participação sócio 4', PCT),
    ('Grupos de clientes que compraram no ano', '0'), ('Grupos que compraram pela 1ª vez no ano', '0'),
]
P = {nome: f'Dados!$AI${5 + i}' for i, (nome, _) in enumerate(PARAM)}
COMP, MES, ANO = P['Competência'], P['Mês'], P['Ano']

LISTAS = [('AK', 'Cliente (grupo)'), ('AL', '1ª compra'), ('AN', 'UF'), ('AO', 'Estado'),
          ('AQ', 'Categoria de despesa'), ('AS', 'Fornecedor do mês')]
DET = [('A', 'Mês', DATA_MES), ('B', 'Rótulo', None), ('C', 'Tipo', None), ('D', 'Item', None),
       ('E', 'Chave', None), ('F', 'Valor', BRL2), ('G', 'Notas', '0')]


def dref(col, row, abs_row=True):
    return f'Dados!${col}${row}' if abs_row else f'Dados!${col}{row}'


def fechado(row):
    return f'Dados!$C${row}=1'


def faixa(col):
    return f'Dados!${col}${DET_INI}:${col}${DET_FIM}'


def det(tipo, item=None, mes=None, de=None, ate=None):
    """SUMIFS no detalhe da aba Dados."""
    partes = [faixa('F'), faixa('C'), f'"{tipo}"']
    if item is not None:
        partes += [faixa('D'), item]
    if mes is not None:
        partes += [faixa('A'), mes]
    if de is not None:
        partes += [faixa('A'), f'">="&{de}']
    if ate is not None:
        partes += [faixa('A'), f'"<="&{ate}']
    return f'SUMIFS({",".join(partes)})'


def vmes(col, data=COMP):
    """Valor de uma coluna da aba Dados no mês `data` (VLOOKUP pela data)."""
    return f'VLOOKUP({data},Dados!$A${D_INI}:$AF${D_FIM},{IDX[col]},FALSE)'


MES_ANT = f'EOMONTH({COMP},-2)+1'
MESMO_MES_ANO_ANT = f'EOMONTH({COMP},-12)'  # último dia do mesmo mês no ano anterior

# ── Estilos ──────────────────────────────────────────────────────────────────
fino = Side(style='thin', color=REGRA)
forte = Side(style='medium', color=VINHO)


def fonte(**kw):
    base = dict(name=FONTE, size=9, color=TINTA)
    base.update(kw)
    return Font(**base)


def fill(cor):
    return PatternFill('solid', start_color=cor, end_color=cor)


ESTILOS = {
    'titulo': dict(font=fonte(size=16, bold=True, color=VINHO)),
    'sub': dict(font=fonte(size=9, color=APAGADO)),
    'cab': dict(font=fonte(bold=True, color=BRANCO), fill=fill(VINHO), alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)),
    'cab_esq': dict(font=fonte(bold=True, color=BRANCO), fill=fill(VINHO), alignment=Alignment(horizontal='left', vertical='center')),
    'rot': dict(font=fonte(), border=Border(bottom=fino), alignment=Alignment(indent=1, vertical='center')),
    'rot_sub': dict(font=fonte(color=APAGADO), border=Border(bottom=fino), alignment=Alignment(indent=2, vertical='center')),
    'num': dict(font=fonte(), border=Border(bottom=fino), number_format=BRL, alignment=Alignment(horizontal='right', vertical='center')),
    'num_sub': dict(font=fonte(color=APAGADO), border=Border(bottom=fino), number_format=BRL, alignment=Alignment(horizontal='right')),
    'pct': dict(font=fonte(color=APAGADO), border=Border(bottom=fino), number_format=PCT, alignment=Alignment(horizontal='right')),
    'int': dict(font=fonte(), border=Border(bottom=fino), number_format=INT, alignment=Alignment(horizontal='right')),
    'tot_rot': dict(font=fonte(bold=True, color=VINHO), fill=fill(CREME), border=Border(top=forte, bottom=fino), alignment=Alignment(indent=1)),
    'tot': dict(font=fonte(bold=True, color=VINHO), fill=fill(CREME), border=Border(top=forte, bottom=fino), number_format=BRL, alignment=Alignment(horizontal='right')),
    'tot_pct': dict(font=fonte(bold=True, color=VINHO), fill=fill(CREME), border=Border(top=forte, bottom=fino), number_format=PCT, alignment=Alignment(horizontal='right')),
    'secao': dict(font=fonte(bold=True, color=TERRA, size=10), alignment=Alignment(vertical='bottom')),
    'texto': dict(font=fonte(), border=Border(bottom=fino), alignment=Alignment(horizontal='center')),
    'nota': dict(font=fonte(size=8, italic=True, color=APAGADO), alignment=Alignment(wrap_text=False)),
    'card_rot': dict(font=fonte(size=8, bold=True, color=APAGADO), fill=fill(CREME), border=Border(top=Side(style='thick', color=VINHO))),
    'card_val': dict(font=fonte(size=18, bold=True, color=VINHO), fill=fill(CREME), number_format=BRL, alignment=Alignment(horizontal='left')),
    'card_det': dict(font=fonte(size=8, color=APAGADO), fill=fill(CREME), alignment=Alignment(horizontal='left')),
    'card_det_num': dict(font=fonte(size=8, bold=True, color=TINTA), fill=fill(CREME), number_format=BRL, alignment=Alignment(horizontal='right')),
    'card_det_pct': dict(font=fonte(size=8, bold=True, color=TINTA), fill=fill(CREME), number_format=PCT, alignment=Alignment(horizontal='right')),
    'link': dict(font=fonte(color=VINHO, bold=True, underline='single'), border=Border(bottom=fino), alignment=Alignment(indent=1)),
    'dados_cab': dict(font=fonte(bold=True, color=TINTA), fill=fill(BEGE), border=Border(bottom=Side(style='thin', color=APAGADO)), alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)),
    'dados': dict(font=fonte()),
}


def put(ws, ref, valor=None, estilo=None, fmt=None):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    if estilo:
        for k, v in ESTILOS[estilo].items():
            setattr(c, k, v)
    if fmt:
        c.number_format = fmt
    return c


def f(expr):
    return '=' + expr


# ── Estrutura comum ──────────────────────────────────────────────────────────

def nova_aba(wb, nome, titulo, subtitulo, cor, largura_rotulo=32, logo_col=None):
    ws = wb.create_sheet(nome)
    ws.sheet_properties.tabColor = cor
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = largura_rotulo
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 8
    ws.row_dimensions[4].height = 20
    put(ws, 'B1', titulo, 'titulo')
    put(ws, 'B2', subtitulo, 'sub')
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = 'AviZee · Workbook de Fechamento'
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.font = FONTE
    ws.oddFooter.right.text = '&A · página &P de &N'
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.font = FONTE
    if logo_col:
        logo(ws, f'{logo_col}1', 26)
    return ws


def logo(ws, ancora, altura):
    img = Image(str(LOGO))
    prop = img.width / img.height
    img.height = altura
    img.width = round(altura * prop)
    ws.add_image(img, ancora)


def cab_meses(ws, linha, col_ini=3, rotulo='R$'):
    put(ws, f'B{linha}', rotulo, 'cab_esq')
    for i, m in enumerate(MESES):
        put(ws, f'{L(col_ini + i)}{linha}', m, 'cab')
    # Mês da competência em destaque (cabeçalho em terracota).
    ws.conditional_formatting.add(
        f'{L(col_ini)}{linha}:{L(col_ini + 11)}{linha}',
        FormulaRule(formula=[f'COLUMN({L(col_ini)}{linha})-{col_ini - 1}={MES}'], fill=fill(TERRA), font=Font(color=BRANCO, bold=True)),
    )


def destacar_mes(ws, faixa_ref, col_ini=3):
    """Fundo creme na coluna do mês da competência."""
    primeira = faixa_ref.split(':')[0]
    col = ''.join(ch for ch in primeira if ch.isalpha())
    lin = ''.join(ch for ch in primeira if ch.isdigit())
    ws.conditional_formatting.add(
        faixa_ref,
        FormulaRule(formula=[f'COLUMN({col}{lin})-{col_ini - 1}={MES}'], fill=fill(CREME)),
    )


def larguras(ws, cols, w):
    for c in cols:
        ws.column_dimensions[c].width = w


def mes_cols(col_ini=3):
    return [L(col_ini + i) for i in range(12)]


def linha_mensal(ws, linha, rotulo, expr_mes, estilo='num', estilo_rot='rot', col_ini=3):
    """expr_mes(i, row_dados_ano, col) -> expressão; i = 0..11."""
    put(ws, f'B{linha}', rotulo, estilo_rot)
    for i in range(12):
        col = L(col_ini + i)
        put(ws, f'{col}{linha}', f(expr_mes(i, JAN[0] + i, col)), estilo)


def nota(ws, ref, texto):
    put(ws, ref, texto, 'nota')


# ── Gráficos ─────────────────────────────────────────────────────────────────

def _txt(tamanho=800, cor=APAGADO, negrito=False):
    cp = CharacterProperties(sz=tamanho, b=negrito, solidFill=cor, latin=DFont(typeface=FONTE))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def estilo_grafico(ch, titulo=None, legenda=True):
    ch.style = 2
    if titulo:
        ch.title = titulo
        ch.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1000, b=True, solidFill=VINHO, latin=DFont(typeface=FONTE)))
        # Sem "overlay" explícito o Excel desenha o título por cima do gráfico.
        ch.title.overlay = False
    if legenda:
        ch.legend.position = 'b'
        ch.legend.overlay = False
        ch.legend.txPr = _txt()
    else:
        ch.legend = None
    for eixo in (getattr(ch, 'x_axis', None), getattr(ch, 'y_axis', None)):
        if eixo is None:
            continue
        eixo.txPr = _txt()
        eixo.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=REGRA))
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.number_format = '#,##0'
    ch.x_axis.majorTickMark = 'none'
    ch.y_axis.majorTickMark = 'none'
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=BEGE))
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))


def serie(ws, linha, c1, c2, titulo_ref, cor, tipo='bar'):
    vals = Reference(ws, min_col=c1, max_col=c2, min_row=linha, max_row=linha)
    s = Series(vals, title=None)
    s.tx = SeriesLabel(strRef=StrRef(f"'{ws.title}'!{titulo_ref}"))
    if tipo == 'bar':
        s.graphicalProperties.solidFill = cor
        s.graphicalProperties.line.noFill = True
        # Sem isso o Excel pinta de branco as barras negativas (pago).
        s.invertIfNegative = False
    else:
        s.graphicalProperties.line.solidFill = cor
        s.graphicalProperties.line.width = 28000
        s.smooth = False
        s.marker.symbol = 'circle'
        s.marker.size = 5
        s.marker.graphicalProperties = GraphicalProperties(solidFill=cor, ln=LineProperties(solidFill=cor))
    return s


def serie_coluna(ws, col, r1, r2, titulo_ref, cor):
    vals = Reference(ws, min_col=col, max_col=col, min_row=r1, max_row=r2)
    s = Series(vals, title=None)
    s.tx = SeriesLabel(strRef=StrRef(f"'{ws.title}'!{titulo_ref}"))
    s.graphicalProperties.solidFill = cor
    s.graphicalProperties.line.noFill = True
    s.invertIfNegative = False
    return s


def rotulos(ch, ws, col1, lin1, col2, lin2):
    """Categorias como texto (strRef): os rótulos são "jan", "set/26", nomes."""
    ref = f"'{ws.title}'!${L(col1)}${lin1}:${L(col2)}${lin2}"
    for s in ch.series:
        s.cat = AxDataSource(strRef=StrRef(f=ref))


def barras_horizontais(ch):
    """Ranking: valor escrito na ponta da barra, sem eixo de valores (que ia
    para o topo e encostava no título)."""
    ch.x_axis.scaling.orientation = 'maxMin'
    ch.y_axis.delete = True
    ch.y_axis.majorGridlines = None
    for s in ch.series:
        s.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False)
        s.dLbls.numFmt = '#,##0'
        s.dLbls.position = 'outEnd'
        s.dLbls.txPr = _txt(800, TINTA)


LINHA_AUX = 90


def auxiliar(ws, origem, n=0):
    """Linha oculta com #N/A nos meses sem número: a linha do gráfico para no
    último mês fechado em vez de cair a zero (texto vira zero no gráfico)."""
    r = LINHA_AUX + n
    put(ws, f'B{r}', f'Auxiliar do gráfico (linha {origem})', 'nota')
    for i in range(12):
        c = L(3 + i)
        put(ws, f'{c}{r}', f(f'IF(ISNUMBER({c}{origem}),{c}{origem},NA())'), 'num_sub')
    ws.row_dimensions[r].hidden = True
    return r


def mostrar_ocultas(ch):
    ch.visible_cells_only = False
    ch.display_blanks = 'gap'


# ── Abas ─────────────────────────────────────────────────────────────────────

def aba_resumo(wb):
    ws = wb.active
    ws.title = 'Resumo'
    ws.sheet_properties.tabColor = VINHO
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2
    larguras(ws, [L(c) for c in range(2, 10)], 15)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.oddFooter.right.text = '&A · página &P de &N'
    ws.oddFooter.right.size = 8
    ws.row_dimensions[1].height = 44
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 16
    logo(ws, 'B1', 46)
    put(ws, 'E1', 'Workbook de Fechamento', 'titulo').alignment = Alignment(vertical='center')
    put(ws, 'E2', f('"Fechamento de "&' + P['Período']), 'secao').font = fonte(size=12, bold=True, color=TERRA)
    # Sem TEXT(): data-base e geração ficam em células com formato de data.
    put(ws, 'E3', 'Data-base', 'sub')
    put(ws, 'F3', f(P['Data-base']), 'sub', fmt=DATA).alignment = Alignment(horizontal='left')
    put(ws, 'G3', 'Gerado em', 'sub').alignment = Alignment(horizontal='right')
    put(ws, 'H3', f(P['Gerado em']), 'sub', fmt=DATA).alignment = Alignment(horizontal='left')
    put(ws, 'I3', 'Valores em R$', 'sub')

    cards = [
        ('FATURAMENTO DO MÊS', vmes('F'),
         ('vs mês anterior', f'IFERROR({vmes("F")}/{vmes("F", MES_ANT)}-1,"")', 'pct'),
         ('notas de saída', P['Notas de saída no mês'], 'int')),
        ('FATURAMENTO NO ANO', f'SUM(Dados!$F${JAN[0]}:$F${JAN[0] + 11})',
         ('do limite do ano', f'IFERROR(SUM(Dados!$F${JAN[0]}:$F${JAN[0] + 11})/{P["Limite de faturamento (ano)"]},"")', 'pct'),
         ('vs mesmo período', f'IFERROR(SUM(Dados!$F${JAN[0]}:$F${JAN[0] + 11})/SUMIFS(Dados!$F${JAN[-1]}:$F${JAN[-1] + 11},Dados!$A${JAN[-1]}:$A${JAN[-1] + 11},"<="&{MESMO_MES_ANO_ANT})-1,"")', 'pct')),
        ('RESULTADO DE CAIXA DO MÊS', f'{vmes("D")}-{vmes("E")}',
         ('recebido', vmes('D'), 'num'),
         ('pago', vmes('E'), 'num')),
        ('RESULTADO DE CAIXA NO ANO', f'SUM(Dados!$D${JAN[0]}:$D${JAN[0] + 11})-SUM(Dados!$E${JAN[0]}:$E${JAN[0] + 11})',
         ('mesmo período ano ant.', f'SUMIFS(Dados!$D${JAN[-1]}:$D${JAN[-1] + 11},Dados!$A${JAN[-1]}:$A${JAN[-1] + 11},"<="&{MESMO_MES_ANO_ANT})-SUMIFS(Dados!$E${JAN[-1]}:$E${JAN[-1] + 11},Dados!$A${JAN[-1]}:$A${JAN[-1] + 11},"<="&{MESMO_MES_ANO_ANT})', 'num'),
         ('margem de caixa', f'IFERROR((SUM(Dados!$D${JAN[0]}:$D${JAN[0] + 11})-SUM(Dados!$E${JAN[0]}:$E${JAN[0] + 11}))/SUM(Dados!$D${JAN[0]}:$D${JAN[0] + 11}),"")', 'pct')),
        ('CAIXA NO FIM DO MÊS', vmes('G'),
         ('disponível', f'{vmes("G")}-{vmes("J")}', 'num'),
         ('bloqueado em garantia', vmes('J'), 'num')),
        ('A RECEBER', '+'.join(vmes(c) for c in 'MNOPQ'),
         ('vencido', f'{vmes("P")}+{vmes("Q")}', 'num'),
         ('vence em até 30 dias', vmes('M'), 'num')),
        ('A PAGAR', '+'.join(vmes(c) for c in 'RSTUV'),
         ('vencido', f'{vmes("U")}+{vmes("V")}', 'num'),
         ('vence em até 30 dias', vmes('R'), 'num')),
        ('ESTOQUE', f'{vmes("K")}+{vmes("L")}',
         ('produtos', vmes('K'), 'num'),
         ('materiais', vmes('L'), 'num')),
    ]
    for k, (rot, valor, d1, d2) in enumerate(cards):
        r0 = 5 if k < 4 else 11
        c0 = 2 + (k % 4) * 2
        a, b = L(c0), L(c0 + 1)
        ws.merge_cells(f'{a}{r0}:{b}{r0}')
        ws.merge_cells(f'{a}{r0 + 1}:{b}{r0 + 1}')
        put(ws, f'{a}{r0}', rot, 'card_rot')
        put(ws, f'{b}{r0}', None, 'card_rot')
        put(ws, f'{a}{r0 + 1}', f(valor), 'card_val')
        put(ws, f'{b}{r0 + 1}', None, 'card_val')
        for j, (dr, dv, tipo) in enumerate((d1, d2)):
            r = r0 + 2 + j
            put(ws, f'{a}{r}', dr, 'card_det')
            put(ws, f'{b}{r}', f(dv), 'card_det_pct' if tipo == 'pct' else 'card_det_num', fmt=INT if tipo == 'int' else None)
        # espaço inferior do cartão
        put(ws, f'{a}{r0 + 4}', None, 'card_det')
        put(ws, f'{b}{r0 + 4}', None, 'card_det')
        ws.row_dimensions[r0 + 1].height = 28
        # separação visual entre cartões
        for r in range(r0, r0 + 5):
            ws[f'{b}{r}'].border = Border(right=Side(style='thick', color=BRANCO), top=ws[f'{b}{r}'].border.top)
    ws.row_dimensions[10].height = 10

    # Conferências
    r = 17
    put(ws, f'B{r}', 'Conferências', 'secao')
    r += 1
    ws.merge_cells(f'B{r}:E{r}')
    for col, txt in zip('BFGHI', ['Conferência', 'Situação', 'Valor', 'Compara com', 'Diferença']):
        put(ws, f'{col}{r}', txt, 'cab' if col != 'B' else 'cab_esq')
    for col in 'CDE':
        put(ws, f'{col}{r}', None, 'cab')
    confs = [
        ('Faturamento por cliente = faturamento do mês', det('Cliente', mes=COMP), vmes('F')),
        ('Faturamento por estado = faturamento do mês', det('Estado', mes=COMP), vmes('F')),
        ('Despesas por categoria = pago no mês', det('Despesa', mes=COMP), vmes('E')),
        # O extrato informado manda: a diferença para o ERP vira o ajuste do mês
        # (informação, não erro; a ponte abaixo confere que ele fecha).
        ('Ajuste ao extrato: caixa informado − calculado pelo ERP', vmes('G'), vmes('H'), 'Ajustado'),
        ('Ponte do caixa: inicial + recebido − pago ± ajuste = final',
         f'{vmes("G", MES_ANT)}+{vmes("D")}-{vmes("E")}+{vmes("I")}', vmes('G')),
    ]
    for desc, va, vb, *alerta in confs:
        r += 1
        ws.merge_cells(f'B{r}:E{r}')
        put(ws, f'B{r}', desc, 'rot')
        put(ws, f'G{r}', f(va), 'num')
        put(ws, f'H{r}', f(vb), 'num')
        put(ws, f'I{r}', f(f'ROUND(G{r}-H{r},2)'), 'num')
        # Conferências toleram arredondamento (< R$ 1); o ajuste aparece a partir de 1 centavo.
        teste = f'I{r}=0' if alerta else f'ABS(I{r})<1'
        put(ws, f'F{r}', f(f'IF({teste},"OK","{alerta[0] if alerta else "Verificar"}")'), 'texto')
    r += 1
    ws.merge_cells(f'B{r}:E{r}')
    put(ws, f'B{r}', 'Extrato do mês conciliado (linhas pendentes)', 'rot')
    put(ws, f'G{r}', f(P['Linhas do extrato pendentes']), 'int')
    put(ws, f'H{r}', None, 'num')
    put(ws, f'I{r}', None, 'num')
    put(ws, f'F{r}', f(f'IF(G{r}=0,"OK","Verificar")'), 'texto')
    r += 1
    ws.merge_cells(f'B{r}:E{r}')
    put(ws, f'B{r}', 'Seguidores informados no mês (LinkedIn · Instagram)', 'rot')
    put(ws, f'G{r}', f(f'IF(ISNUMBER({vmes("W")}),{vmes("W")},"sem dado")'), 'int')
    put(ws, f'H{r}', f(f'IF(ISNUMBER({vmes("X")}),{vmes("X")},"sem dado")'), 'int')
    put(ws, f'I{r}', None, 'num')
    put(ws, f'F{r}', f(f'IF(AND(ISNUMBER(G{r}),ISNUMBER(H{r})),"OK","Verificar")'), 'texto')
    conf_fim = r
    ws.conditional_formatting.add(f'F19:F{conf_fim}', FormulaRule(formula=['F19="OK"'], font=Font(color=BOM, bold=True)))
    ws.conditional_formatting.add(f'F19:F{conf_fim}', FormulaRule(formula=['F19="Verificar"'], fill=fill(TERRA), font=Font(color=BRANCO, bold=True)))
    ws.conditional_formatting.add(f'F19:F{conf_fim}', FormulaRule(formula=['F19="Ajustado"'], fill=fill(CREME), font=Font(color=TERRA, bold=True)))

    # Índice
    r = conf_fim + 2
    put(ws, f'B{r}', 'Abas', 'secao')
    indice = [
        ('Resultado', 'Recebido, pago e resultado de caixa, por mês, trimestre e ano, contra o ano anterior'),
        ('Faturamento', 'Notas de saída em três anos, acumulado contra a meta e o limite do ano'),
        ('Caixa', 'Ponte do caixa mês a mês (com rendimentos), bloqueado, disponível e meses de despesa cobertos'),
        ('Despesas', 'Pagamentos por categoria e os maiores fornecedores do mês'),
        ('Clientes', 'Faturamento por grupo econômico, participação, curva ABC e 1ª compra'),
        ('Estados', 'Faturamento por estado e trimestre, contra o mesmo período do ano anterior'),
        ('Capital de giro', 'A receber e a pagar por faixa de vencimento, nos últimos 12 meses'),
        ('Estoque', 'Produtos e materiais em 24 meses e cobertura em meses de faturamento'),
        ('Sócios', 'Retiradas do ano e resultado atribuído pela participação'),
        ('Redes', 'Seguidores no LinkedIn e no Instagram, mês a mês'),
        ('Dados', 'Base de todas as abas: um mês por linha e o detalhe do mês (com centavos)'),
    ]
    for nome, desc in indice:
        r += 1
        c = put(ws, f'B{r}', nome, 'link')
        c.hyperlink = Hyperlink(ref=f'B{r}', location=f"'{nome}'!A1", display=nome)
        ws.merge_cells(f'C{r}:I{r}')
        put(ws, f'C{r}', desc, 'rot')
        for col in 'DEFGHI':
            put(ws, f'{col}{r}', None, 'rot')
    r += 2
    nota(ws, f'B{r}', 'Valores em R$ sem centavos; negativos em vermelho e entre parênteses; "–" é zero. Os centavos estão na aba Dados.')
    nota(ws, f'B{r + 1}', 'Regime de caixa: recebimentos e pagamentos do mês, sem transferências entre contas. Faturamento: notas de saída emitidas.')
    ws.print_area = f'A1:I{r + 1}'
    return ws


def aba_resultado(wb):
    ws = nova_aba(wb, 'Resultado', f('"Resultado de caixa · "&' + ANO),
                  'Regime de caixa: o que entrou e saiu das contas no mês, sem transferências entre contas · R$',
                  VINHO, logo_col='R')
    larguras(ws, mes_cols(), 10)
    larguras(ws, ['O', 'P', 'Q', 'R'], 10.5)
    ws.column_dimensions['S'].width = 12
    cab_meses(ws, 4)
    for i, c in enumerate('OPQR'):
        put(ws, f'{c}4', f'{i + 1}º tri', 'cab')
    put(ws, 'S4', f(ANO), 'cab', fmt='0')

    def g(row, expr):
        return f'IF({fechado(row)},{expr},"")'

    linha_mensal(ws, 5, 'Recebido', lambda i, r, c: g(r, dref('D', r)))
    linha_mensal(ws, 6, 'Pago', lambda i, r, c: g(r, '-' + dref('E', r)))
    linha_mensal(ws, 7, 'Resultado de caixa', lambda i, r, c: f'IF({c}5="","",{c}5+{c}6)', 'tot', 'tot_rot')
    linha_mensal(ws, 8, 'Resultado acumulado no ano',
                 lambda i, r, c: f'IF({c}7="","",{c}7)' if i == 0 else f'IF({c}7="","",N({L(2 + i)}8)+{c}7)', 'num_sub', 'rot_sub')
    linha_mensal(ws, 9, 'Margem de caixa', lambda i, r, c: f'IF({c}5="","",IFERROR({c}7/{c}5,""))', 'pct', 'rot_sub')
    put(ws, 'B11', f('"Ano anterior · "&(' + ANO + '-1)'), 'secao')
    for row, rot, col, sinal in [(12, 'Recebido', 'D', ''), (13, 'Pago', 'E', '-')]:
        put(ws, f'B{row}', rot, 'rot')
        for i in range(12):
            r = JAN[-1] + i
            put(ws, f'{L(3 + i)}{row}', f(f'IF({fechado(r)},{sinal}{dref(col, r)},"")'), 'num')
    linha_mensal(ws, 14, 'Resultado de caixa', lambda i, r, c: f'IF({c}12="","",{c}12+{c}13)', 'tot', 'tot_rot')
    linha_mensal(ws, 15, 'Diferença do resultado vs ano anterior', lambda i, r, c: f'IF({c}7="","",{c}7-{c}14)', 'num_sub', 'rot_sub')

    for row in (5, 6, 7, 12, 13, 14, 15):
        est = 'tot' if row in (7, 14) else ('num_sub' if row == 15 else 'num')
        for q, c in enumerate('OPQR'):
            a, b = L(3 + q * 3), L(5 + q * 3)
            put(ws, f'{c}{row}', f(f'SUM({a}{row}:{b}{row})'), est)
        put(ws, f'S{row}', f(f'SUM(C{row}:N{row})'), est)
    for c in 'OPQRS':
        put(ws, f'{c}8', None, 'num_sub')
        put(ws, f'{c}9', f(f'IFERROR({c}7/{c}5,"")'), 'pct')
    destacar_mes(ws, 'C5:N9')
    ws.freeze_panes = 'C5'
    nota(ws, 'B16', 'Trimestres e ano somam só os meses fechados. Meses futuros ficam vazios.')

    ch = BarChart()
    ch.type = 'col'
    ch.grouping = 'clustered'
    ch.overlap = 100
    ch.gapWidth = 60
    ch.series.append(serie(ws, 5, 3, 14, '$B$5', VINHO))
    ch.series.append(serie(ws, 6, 3, 14, '$B$6', TERRA2))
    rotulos(ch, ws, 3, 4, 14, 4)
    ln = LineChart()
    ln.series.append(serie(ws, auxiliar(ws, 8), 3, 14, '$B$8', TINTA, 'line'))
    rotulos(ln, ws, 3, 4, 14, 4)
    estilo_grafico(ch, 'Recebido × pago e resultado acumulado')
    ch.x_axis.tickLblPos = 'low'
    mostrar_ocultas(ch)
    ch += ln
    ch.width, ch.height = 30, 8.5
    ws.add_chart(ch, 'B18')
    return ws


def aba_faturamento(wb):
    ws = nova_aba(wb, 'Faturamento', f('"Faturamento · "&' + ANO),
                  'Notas fiscais de saída emitidas no mês · coluna Ano: variação contra o mesmo período do ano anterior · R$',
                  TERRA, logo_col='N')
    larguras(ws, mes_cols(), 10)
    ws.column_dimensions['O'].width = 12
    cab_meses(ws, 4)
    put(ws, 'O4', 'Ano', 'cab')
    for row, off in [(5, -2), (6, -1), (7, 0)]:
        put(ws, f'B{row}', f(f'{ANO}-{-off}' if off else ANO), 'tot_rot' if off == 0 else 'rot', fmt='0')
        ws[f'B{row}'].alignment = Alignment(horizontal='left', indent=1)
        for i in range(12):
            r = JAN[off] + i
            put(ws, f'{L(3 + i)}{row}', f(f'IF({fechado(r)},{dref("F", r)},"")'), 'tot' if off == 0 else 'num')
        put(ws, f'O{row}', f(f'SUM(C{row}:N{row})'), 'tot' if off == 0 else 'num')
    linha_mensal(ws, 8, 'Variação vs ano anterior', lambda i, r, c: f'IF({c}7="","",IFERROR({c}7/{c}6-1,""))', 'pct', 'rot_sub')
    put(ws, 'O8', f(f'IFERROR(O7/SUMIFS(Dados!$F${JAN[-1]}:$F${JAN[-1] + 11},Dados!$A${JAN[-1]}:$A${JAN[-1] + 11},"<="&{MESMO_MES_ANO_ANT})-1,"")'), 'pct')

    put(ws, 'B10', 'Acumulado, meta e limite', 'secao')
    linha_mensal(ws, 11, 'Acumulado no ano', lambda i, r, c: f'IF({c}7="","",{c}7)' if i == 0 else f'IF({c}7="","",N({L(2 + i)}11)+{c}7)', 'tot', 'tot_rot')
    linha_mensal(ws, 12, 'Acumulado do ano anterior', lambda i, r, c: f'N({c}6)' if i == 0 else f'{L(2 + i)}12+N({c}6)')
    linha_mensal(ws, 13, 'Meta acumulada', lambda i, r, c: f'{c}12*(1+{P["Meta de crescimento do ano"]})')
    linha_mensal(ws, 14, 'Realizado / meta', lambda i, r, c: f'IF({c}11="","",IFERROR({c}11/{c}13,""))', 'pct', 'rot_sub')
    linha_mensal(ws, 15, 'Limite de faturamento do ano', lambda i, r, c: P['Limite de faturamento (ano)'], 'num_sub', 'rot_sub')
    linha_mensal(ws, 16, 'Uso do limite', lambda i, r, c: f'IF({c}11="","",IFERROR({c}11/{c}15,""))', 'pct', 'rot_sub')
    put(ws, 'O11', f('N(O7)'), 'tot')
    put(ws, 'O12', f('O6'), 'num')
    put(ws, 'O13', f('O12*(1+' + P['Meta de crescimento do ano'] + ')'), 'num')
    put(ws, 'O14', f('IFERROR(O11/O13,"")'), 'pct')
    put(ws, 'O15', f(P['Limite de faturamento (ano)']), 'num_sub')
    put(ws, 'O16', f('IFERROR(O11/O15,"")'), 'pct')
    nota(ws, 'B17', 'Meta: faturamento do ano anterior acrescido da meta de crescimento (aba Dados, parâmetros). Limite: teto de faturamento do ano.')
    destacar_mes(ws, 'C5:N16')
    ws.freeze_panes = 'C5'

    ch = BarChart()
    ch.type = 'col'
    ch.gapWidth = 50
    ch.series.append(serie(ws, 6, 3, 14, '$B$6', BEGE))
    ch.series.append(serie(ws, 7, 3, 14, '$B$7', VINHO))
    rotulos(ch, ws, 3, 4, 14, 4)
    estilo_grafico(ch, 'Faturamento por mês')
    ch.width, ch.height = 15, 8
    ws.add_chart(ch, 'B19')

    ln = LineChart()
    ln.series.append(serie(ws, auxiliar(ws, 11), 3, 14, '$B$11', VINHO, 'line'))
    ln.series.append(serie(ws, 13, 3, 14, '$B$13', TERRA2, 'line'))
    ln.series.append(serie(ws, 15, 3, 14, '$B$15', BEGE, 'line'))
    ln.series[2].marker.symbol = 'none'
    ln.series[2].graphicalProperties.line.dashStyle = 'dash'
    rotulos(ln, ws, 3, 4, 14, 4)
    estilo_grafico(ln, 'Acumulado × meta × limite')
    mostrar_ocultas(ln)
    ln.width, ln.height = 15, 8
    ws.add_chart(ln, 'H19')
    return ws


def aba_caixa(wb):
    ws = nova_aba(wb, 'Caixa', f('"Caixa · "&' + ANO),
                  'Do caixa inicial ao final; o ajuste ao extrato (informado − calculado pelo ERP) fica numa linha própria · R$',
                  VINHO, logo_col='N')
    larguras(ws, mes_cols(), 10.5)
    cab_meses(ws, 4)

    def g(r, expr):
        return f'IF({fechado(r)},{expr},"")'

    linha_mensal(ws, 5, 'Caixa inicial', lambda i, r, c: g(r, dref('G', r - 1)), 'num_sub', 'rot_sub')
    # O recebido da aba Dados inclui os rendimentos; aqui eles ficam numa linha própria.
    linha_mensal(ws, 6, '+ Recebimentos', lambda i, r, c: g(r, f'{dref("D", r)}-{dref("AG", r)}'))
    linha_mensal(ws, 7, '+ Rendimentos de aplicações', lambda i, r, c: g(r, dref('AG', r)))
    linha_mensal(ws, 8, '− Pago', lambda i, r, c: g(r, '-' + dref('E', r)))
    linha_mensal(ws, 9, '± Ajuste ao extrato', lambda i, r, c: g(r, dref('I', r)))
    linha_mensal(ws, 10, 'Caixa final', lambda i, r, c: g(r, dref('G', r)), 'tot', 'tot_rot')
    linha_mensal(ws, 11, 'bloqueado em garantia', lambda i, r, c: g(r, dref('J', r)), 'num_sub', 'rot_sub')
    linha_mensal(ws, 12, 'disponível', lambda i, r, c: g(r, f'{dref("G", r)}-{dref("J", r)}'))
    ws['B12'].font = fonte(bold=True)
    linha_mensal(ws, 13, 'Meses de despesa cobertos', lambda i, r, c: g(r, f'IFERROR({c}12/AVERAGE(Dados!$E${r - 2}:$E${r}),"")'), 'num_sub', 'rot_sub')
    for c in mes_cols():
        ws[f'{c}13'].number_format = MESES_TXT
    linha_mensal(ws, 14, 'Caixa final do ano anterior', lambda i, r, c: f'IF({fechado(r - 12)},{dref("G", r - 12)},"")', 'num_sub', 'rot_sub')
    linha_mensal(ws, 15, 'Diferença na ponte (deve ser zero)', lambda i, r, c: f'IF({c}10="","",ROUND({c}5+{c}6+{c}7+{c}8+{c}9-{c}10,2))', 'num_sub', 'rot_sub')
    destacar_mes(ws, 'C5:N15')
    ws.freeze_panes = 'C5'
    nota(ws, 'B16', 'Meses de despesa cobertos: caixa disponível ÷ média do pago nos últimos 3 meses. Rendimentos: CDB, reserva remunerada e CDI informados no fechamento.')

    ch = BarChart()
    ch.type = 'col'
    ch.grouping = 'stacked'
    ch.overlap = 100
    ch.gapWidth = 50
    ch.series.append(serie(ws, 12, 3, 14, '$B$12', VINHO))
    ch.series.append(serie(ws, 11, 3, 14, '$B$11', TERRA3))
    rotulos(ch, ws, 3, 4, 14, 4)
    ln = LineChart()
    ln.series.append(serie(ws, 14, 3, 14, '$B$14', APAGADO, 'line'))
    ln.series[0].graphicalProperties.line.dashStyle = 'dash'
    rotulos(ln, ws, 3, 4, 14, 4)
    estilo_grafico(ch, 'Caixa no fim do mês: disponível e bloqueado')
    ch += ln
    ch.width, ch.height = 26, 8.5
    ws.add_chart(ch, 'B18')
    return ws


def aba_despesas(wb):
    ws = nova_aba(wb, 'Despesas', f('"Despesas por categoria · "&' + ANO),
                  'Pagamentos do mês por categoria (mesma base do pago em Resultado) e os maiores fornecedores pagos no mês · R$',
                  VINHO, logo_col='P')
    larguras(ws, mes_cols(), 9.5)
    larguras(ws, ['O', 'Q'], 11.5)
    ws.column_dimensions['P'].width = 8
    cab_meses(ws, 4, rotulo='Categoria')
    put(ws, 'O4', 'Ano', 'cab')
    put(ws, 'P4', '% do ano', 'cab')
    put(ws, 'Q4', f(f'"Em "&({ANO}-1)'), 'cab')
    r0, n = 5, 25
    tot = r0 + n
    for k in range(n):
        r = r0 + k
        lista = f'Dados!$AQ${5 + k}'
        put(ws, f'B{r}', f(f'IF({lista}="","",{lista})'), 'rot')
        for i in range(12):
            c = L(3 + i)
            rd = JAN[0] + i
            put(ws, f'{c}{r}', f(f'IF($B{r}="","",IF({fechado(rd)},{det("Despesa", f"$B{r}", dref("A", rd))},""))'), 'num')
        put(ws, f'O{r}', f(f'IF($B{r}="","",SUM(C{r}:N{r}))'), 'num')
        put(ws, f'P{r}', f(f'IF($B{r}="","",IFERROR(O{r}/$O${tot},""))'), 'pct')
        put(ws, f'Q{r}', f(f'IF($B{r}="","",{det("Despesa", f"$B{r}", de=dref("A", JAN[-1]), ate=dref("A", JAN[-1] + 11))})'), 'num_sub')
    put(ws, f'B{tot}', 'Total pago', 'tot_rot')
    for i in range(12):
        c = L(3 + i)
        put(ws, f'{c}{tot}', f(f'IF({fechado(JAN[0] + i)},SUM({c}{r0}:{c}{tot - 1}),"")'), 'tot')
    put(ws, f'O{tot}', f(f'SUM(O{r0}:O{tot - 1})'), 'tot')
    put(ws, f'P{tot}', f(f'IFERROR(O{tot}/O{tot},"")'), 'tot_pct')
    put(ws, f'Q{tot}', f(f'IF(SUM(Q{r0}:Q{tot - 1})=0,"sem detalhe",SUM(Q{r0}:Q{tot - 1}))'), 'tot')
    linha_mensal(ws, tot + 1, 'Pago no mês (Resultado)', lambda i, r, c: f'IF({fechado(r)},{dref("E", r)},"")', 'num_sub', 'rot_sub')
    put(ws, f'O{tot + 1}', f(f'SUM(C{tot + 1}:N{tot + 1})'), 'num_sub')
    destacar_mes(ws, f'C{r0}:N{tot}')
    ws.freeze_panes = 'C5'

    rf = tot + 4
    put(ws, f'B{rf - 1}', f('"Maiores fornecedores pagos em "&' + P['Período']), 'secao')
    put(ws, f'B{rf}', 'Fornecedor', 'cab_esq')
    put(ws, f'C{rf}', 'Valor', 'cab')
    put(ws, f'D{rf}', '% do pago', 'cab')
    for k in range(10):
        r = rf + 1 + k
        lista = f'Dados!$AS${5 + k}'
        put(ws, f'B{r}', f(f'IF({lista}="","",{lista})'), 'rot')
        # Razão social longa: reduz a fonte em vez de cortar o nome.
        ws[f'B{r}'].alignment = Alignment(indent=1, vertical='center', shrink_to_fit=True)
        put(ws, f'C{r}', f(f'IF($B{r}="","",{det("Fornecedor", f"$B{r}", COMP)})'), 'num')
        put(ws, f'D{r}', f(f'IF($B{r}="","",IFERROR(C{r}/{vmes("E")},""))'), 'pct')
    nota(ws, f'B{rf + 12}', 'Inclui retiradas dos sócios e impostos, como no extrato.')
    nota(ws, f'B{rf + 13}', '"sem detalhe": ano importado só com o total pago.')

    ch = BarChart()
    ch.type = 'bar'
    ch.gapWidth = 40
    ch.series.append(serie_coluna(ws, 15, r0, r0 + 9, '$O$4', VINHO))
    rotulos(ch, ws, 2, r0, 2, r0 + 9)
    estilo_grafico(ch, 'Maiores categorias no ano', legenda=False)
    barras_horizontais(ch)
    ch.width, ch.height = 17, 8
    ws.add_chart(ch, f'F{rf - 1}')
    return ws


def aba_clientes(wb):
    ws = nova_aba(wb, 'Clientes', f('"Clientes · "&' + ANO),
                  'Faturamento por grupo econômico (CNPJ raiz), do maior para o menor no ano · curva ABC: A até 80% do total, B até 95% · R$',
                  TERRA, largura_rotulo=28, logo_col='R')
    larguras(ws, mes_cols(), 9.5)
    ws.column_dimensions['O'].width = 11.5
    ws.column_dimensions['P'].width = 7
    ws.column_dimensions['Q'].width = 7
    ws.column_dimensions['R'].width = 9.5
    ws.column_dimensions['S'].width = 11.5
    cab_meses(ws, 4, rotulo='Cliente (grupo)')
    for c, t in zip('OPQR', ['Ano', '%', 'Curva', '1ª compra']):
        put(ws, f'{c}4', t, 'cab')
    put(ws, 'S4', f(f'"Em "&({ANO}-1)'), 'cab')
    r0, n = 5, SLOTS_LISTA
    outros, tot = r0 + n, r0 + n + 1
    for k in range(n):
        r = r0 + k
        lista = f'Dados!$AK${5 + k}'
        put(ws, f'B{r}', f(f'IF({lista}="","",{lista})'), 'rot')
        for i in range(12):
            rd = JAN[0] + i
            put(ws, f'{L(3 + i)}{r}', f(f'IF($B{r}="","",IF({fechado(rd)},{det("Cliente", f"$B{r}", dref("A", rd))},""))'), 'num')
        put(ws, f'O{r}', f(f'IF($B{r}="","",SUM(C{r}:N{r}))'), 'num')
        put(ws, f'P{r}', f(f'IF($B{r}="","",IFERROR(O{r}/$O${tot},""))'), 'pct')
        put(ws, f'Q{r}', f(f'IF(N(O{r})=0,"",IF(SUM($P${r0}:P{r})-P{r}<0.8,"A",IF(SUM($P${r0}:P{r})-P{r}<0.95,"B","C")))'), 'texto')
        put(ws, f'R{r}', f(f'IF(Dados!$AL${5 + k}="","",Dados!$AL${5 + k})'), 'texto')
        ws[f'R{r}'].font = fonte(color=APAGADO)
        put(ws, f'S{r}', f(f'IF($B{r}="","",{det("Cliente", f"$B{r}", de=dref("A", JAN[-1]), ate=dref("A", JAN[-1] + 11))})'), 'num_sub')
    put(ws, f'B{outros}', 'Outros', 'rot')
    for i in range(12):
        c = L(3 + i)
        put(ws, f'{c}{outros}', f(f'IF({c}{tot}="","",ROUND({c}{tot}-SUM({c}{r0}:{c}{outros - 1}),2))'), 'num')
    put(ws, f'O{outros}', f(f'SUM(C{outros}:N{outros})'), 'num')
    put(ws, f'P{outros}', f(f'IFERROR(O{outros}/$O${tot},"")'), 'pct')
    put(ws, f'S{outros}', f(f'ROUND(S{tot}-SUM(S{r0}:S{outros - 1}),2)'), 'num_sub')
    put(ws, f'B{tot}', 'Total', 'tot_rot')
    for i in range(12):
        rd = JAN[0] + i
        put(ws, f'{L(3 + i)}{tot}', f(f'IF({fechado(rd)},{det("Cliente", mes=dref("A", rd))},"")'), 'tot')
    put(ws, f'O{tot}', f(f'SUM(C{tot}:N{tot})'), 'tot')
    put(ws, f'P{tot}', f(f'IFERROR(O{tot}/O{tot},"")'), 'tot_pct')
    for c in 'QR':
        put(ws, f'{c}{tot}', None, 'tot')
    put(ws, f'S{tot}', f(det('Cliente', de=dref('A', JAN[-1]), ate=dref('A', JAN[-1] + 11))), 'tot')
    ws.conditional_formatting.add(f'Q{r0}:Q{outros - 1}', FormulaRule(formula=[f'Q{r0}="A"'], font=Font(color=VINHO, bold=True)))
    destacar_mes(ws, f'C{r0}:N{tot}')
    ws.freeze_panes = 'C5'
    r = tot + 2
    put(ws, f'B{r}', 'Grupos que compraram no ano', 'rot')
    put(ws, f'C{r}', f(P['Grupos de clientes que compraram no ano']), 'int')
    put(ws, f'B{r + 1}', 'Grupos que compraram pela 1ª vez no ano', 'rot')
    put(ws, f'C{r + 1}', f(P['Grupos que compraram pela 1ª vez no ano']), 'int')

    ch = BarChart()
    ch.type = 'bar'
    ch.gapWidth = 40
    ch.series.append(serie_coluna(ws, 15, r0, r0 + 9, '$O$4', TERRA))
    rotulos(ch, ws, 2, r0, 2, r0 + 9)
    estilo_grafico(ch, 'Dez maiores clientes no ano', legenda=False)
    barras_horizontais(ch)
    ch.width, ch.height = 17, 8
    ws.add_chart(ch, f'E{r + 3}')
    return ws


def aba_estados(wb):
    ws = nova_aba(wb, 'Estados', f('"Faturamento por estado · "&' + ANO),
                  'Estado do cliente na nota de saída, por trimestre · comparação com o mesmo período do ano anterior · R$',
                  TERRA, largura_rotulo=22, logo_col='K')
    ws.column_dimensions['C'].width = 6
    larguras(ws, ['D', 'E', 'F', 'G', 'H', 'J', 'L'], 11.5)
    larguras(ws, ['I', 'K'], 8.5)
    heads = ['UF', '1º tri', '2º tri', '3º tri', '4º tri', 'Ano', '%']
    put(ws, 'B4', 'Estado', 'cab_esq')
    for c, t in zip('CDEFGHI', heads):
        put(ws, f'{c}4', t, 'cab')
    put(ws, 'J4', f(f'"Mesmo período "&({ANO}-1)'), 'cab')
    put(ws, 'K4', 'Variação', 'cab')
    put(ws, 'L4', f(f'"Ano "&({ANO}-1)'), 'cab')
    ws.row_dimensions[4].height = 28
    r0, n = 5, 28
    tot = r0 + n
    for k in range(n):
        r = r0 + k
        put(ws, f'B{r}', f(f'IF(Dados!$AO${5 + k}="","",Dados!$AO${5 + k})'), 'rot')
        put(ws, f'C{r}', f(f'IF(Dados!$AN${5 + k}="","",Dados!$AN${5 + k})'), 'texto')
        for q, c in enumerate('DEFG'):
            de, ate = dref('A', JAN[0] + 3 * q), dref('A', JAN[0] + 3 * q + 2)
            put(ws, f'{c}{r}', f(f'IF($C{r}="","",{det("Estado", f"$C{r}", de=de, ate=ate)})'), 'num')
        put(ws, f'H{r}', f(f'IF($C{r}="","",SUM(D{r}:G{r}))'), 'num')
        put(ws, f'I{r}', f(f'IF($C{r}="","",IFERROR(H{r}/$H${tot},""))'), 'pct')
        put(ws, f'J{r}', f(f'IF($C{r}="","",{det("Estado", f"$C{r}", de=dref("A", JAN[-1]), ate=MESMO_MES_ANO_ANT)})'), 'num_sub')
        put(ws, f'K{r}', f(f'IF($C{r}="","",IFERROR(H{r}/J{r}-1,""))'), 'pct')
        put(ws, f'L{r}', f(f'IF($C{r}="","",{det("Estado", f"$C{r}", de=dref("A", JAN[-1]), ate=dref("A", JAN[-1] + 11))})'), 'num_sub')
    put(ws, f'B{tot}', 'Total', 'tot_rot')
    put(ws, f'C{tot}', None, 'tot')
    for c in 'DEFGHJL':
        put(ws, f'{c}{tot}', f(f'SUM({c}{r0}:{c}{tot - 1})'), 'tot')
    put(ws, f'I{tot}', f(f'IFERROR(H{tot}/H{tot},"")'), 'tot_pct')
    put(ws, f'K{tot}', f(f'IFERROR(H{tot}/J{tot}-1,"")'), 'tot_pct')
    ws.conditional_formatting.add(f'H{r0}:H{tot - 1}', ColorScaleRule(start_type='num', start_value=0, start_color=CREME,
                                                                      mid_type='percentile', mid_value=60, mid_color=TERRA3,
                                                                      end_type='max', end_color=VINHO2))
    ws.freeze_panes = 'D5'
    nota(ws, f'B{tot + 1}', 'S/UF: notas de clientes sem estado no cadastro.')

    ch = BarChart()
    ch.type = 'bar'
    ch.gapWidth = 40
    ch.series.append(serie_coluna(ws, 8, r0, r0 + 9, '$H$4', TERRA))
    rotulos(ch, ws, 2, r0, 2, r0 + 9)
    estilo_grafico(ch, 'Dez maiores estados no ano', legenda=False)
    barras_horizontais(ch)
    ch.width, ch.height = 15, 9
    ws.add_chart(ch, 'N4')
    return ws


def aba_capital(wb):
    ws = nova_aba(wb, 'Capital de giro', f('"Capital de giro · últimos 12 meses até "&' + P['Período']),
                  'Contas a receber e a pagar em aberto no fim de cada mês, por faixa de vencimento · R$',
                  TERRA2, largura_rotulo=30, logo_col='M')
    larguras(ws, mes_cols(), 10.5)
    put(ws, 'B4', 'R$', 'cab_esq')
    for k in range(12):
        put(ws, f'{L(3 + k)}4', f(f'VLOOKUP(EOMONTH({COMP},{k - 12})+1,Dados!$A${D_INI}:$B${D_FIM},2,FALSE)'), 'cab')
    put(ws, 'N4', None, 'cab')
    ws.conditional_formatting.add('N4', FormulaRule(formula=['TRUE'], fill=fill(TERRA), font=Font(color=BRANCO, bold=True)))

    faixas = ['A vencer em até 30 dias', 'A vencer de 31 a 90 dias', 'A vencer em mais de 90 dias',
              'Vencido há até 30 dias', 'Vencido há mais de 30 dias']

    def v(col, c):
        return f'VLOOKUP({c}$4,Dados!$B${D_INI}:$AF${D_FIM},{IDX[col] - 1},FALSE)'

    def bloco(r_sec, titulo, cols):
        put(ws, f'B{r_sec}', titulo, 'secao')
        for j, (nome, col) in enumerate(zip(faixas, cols)):
            r = r_sec + 1 + j
            put(ws, f'B{r}', nome, 'rot')
            for k in range(12):
                c = L(3 + k)
                put(ws, f'{c}{r}', f(v(col, c)), 'num')
            if j >= 3:
                ws[f'B{r}'].font = fonte(color=RUIM)
        r_tot = r_sec + 6
        put(ws, f'B{r_tot}', f'Total {titulo.lower()}', 'tot_rot')
        for k in range(12):
            c = L(3 + k)
            put(ws, f'{c}{r_tot}', f(f'SUM({c}{r_sec + 1}:{c}{r_sec + 5})'), 'tot')
        return r_tot

    tr = bloco(5, 'A receber', 'MNOPQ')
    tp = bloco(13, 'A pagar', 'RSTUV')
    put(ws, 'B21', 'Saldo (a receber − a pagar)', 'tot_rot')
    for k in range(12):
        c = L(3 + k)
        put(ws, f'{c}21', f(f'{c}{tr}-{c}{tp}'), 'tot')
    put(ws, 'B22', 'Vencido / total a receber', 'rot_sub')
    put(ws, 'B23', 'Vencido / total a pagar', 'rot_sub')
    for k in range(12):
        c = L(3 + k)
        put(ws, f'{c}22', f(f'IFERROR(({c}9+{c}10)/{c}{tr},"")'), 'pct')
        put(ws, f'{c}23', f(f'IFERROR(({c}17+{c}18)/{c}{tp},"")'), 'pct')
    ws.conditional_formatting.add('N5:N23', FormulaRule(formula=['TRUE'], fill=fill(CREME)))
    ws.freeze_panes = 'C5'

    ch = BarChart()
    ch.type = 'col'
    ch.gapWidth = 60
    ch.series.append(serie(ws, tr, 3, 14, f'$B${tr}', TERRA))
    ch.series.append(serie(ws, tp, 3, 14, f'$B${tp}', VINHO))
    rotulos(ch, ws, 3, 4, 14, 4)
    estilo_grafico(ch, 'A receber × a pagar no fim do mês')
    ch.width, ch.height = 26, 8
    ws.add_chart(ch, 'B25')
    return ws


def aba_estoque(wb):
    ws = nova_aba(wb, 'Estoque', f('"Estoque · 24 meses até "&' + P['Período']),
                  'Saldo no fim do mês pelo custo · cobertura: estoque ÷ faturamento médio dos últimos 3 meses · R$',
                  TERRA2, largura_rotulo=10, logo_col='P')
    larguras(ws, ['C', 'D', 'E', 'F', 'G'], 12.5)
    ws.column_dimensions['H'].width = 12
    for c, t in zip('BCDEFGH', ['Mês', 'Produtos', 'Materiais', 'Total', 'Variação no mês', 'Faturamento do mês', 'Cobertura (meses)']):
        put(ws, f'{c}4', t, 'cab' if c != 'B' else 'cab_esq')
    ws.row_dimensions[4].height = 28
    r0 = 5
    for k in range(24):
        r = r0 + k
        put(ws, f'B{r}', f(f'VLOOKUP(EOMONTH({COMP},{k - 24})+1,Dados!$A${D_INI}:$B${D_FIM},2,FALSE)'), 'rot')

        def v(col):
            return f'VLOOKUP($B{r},Dados!$B${D_INI}:$AF${D_FIM},{IDX[col] - 1},FALSE)'
        put(ws, f'C{r}', f(v('K')), 'num')
        put(ws, f'D{r}', f(v('L')), 'num')
        put(ws, f'E{r}', f(f'C{r}+D{r}'), 'num')
        ws[f'E{r}'].font = fonte(bold=True)
        put(ws, f'F{r}', f('""' if k == 0 else f'E{r}-E{r - 1}'), 'num_sub')
        put(ws, f'G{r}', f(v('F')), 'num_sub')
        media = f'AVERAGE(G{max(r0, r - 2)}:G{r})'
        put(ws, f'H{r}', f(f'IFERROR(E{r}/{media},"")'), 'num_sub', fmt=MESES_TXT)
    ultima = r0 + 23
    for c in 'BCDEFGH':
        ws[f'{c}{ultima}'].fill = fill(CREME)
        ws[f'{c}{ultima}'].font = fonte(bold=True, color=VINHO)
    ws[f'H{ultima}'].number_format = MESES_TXT
    ws.freeze_panes = 'C5'

    ch = BarChart()
    ch.type = 'col'
    ch.grouping = 'stacked'
    ch.overlap = 100
    ch.gapWidth = 40
    ch.series.append(serie_coluna(ws, 3, r0, ultima, '$C$4', VINHO))
    ch.series.append(serie_coluna(ws, 4, r0, ultima, '$D$4', TERRA3))
    rotulos(ch, ws, 2, r0, 2, ultima)
    estilo_grafico(ch, 'Estoque no fim do mês')
    ch.width, ch.height = 18, 10
    ws.add_chart(ch, 'J4')
    return ws


def aba_socios(wb):
    ws = nova_aba(wb, 'Sócios', f('"Sócios · "&' + ANO),
                  'Retiradas (pró-labore) aprovadas no mês · resultado atribuído: resultado do ano até o mês × participação · R$',
                  VINHO, largura_rotulo=26, logo_col='O')
    larguras(ws, mes_cols(), 9.5)
    ws.column_dimensions['O'].width = 11.5
    ws.column_dimensions['P'].width = 11
    cab_meses(ws, 4, rotulo='Sócio')
    put(ws, 'O4', 'Ano', 'cab')
    put(ws, 'P4', 'Participação', 'cab')
    put(ws, 'B5', 'Retiradas', 'secao')
    for k in range(4):
        r = 6 + k
        nome = P[f'Sócio {k + 1}']
        put(ws, f'B{r}', f(f'IF({nome}="","",{nome})'), 'rot')
        col = L(25 + k)  # Y..AB
        for i in range(12):
            rd = JAN[0] + i
            put(ws, f'{L(3 + i)}{r}', f(f'IF(OR($B{r}="",Dados!$C${rd}<>1),"",{dref(col, rd)})'), 'num')
        put(ws, f'O{r}', f(f'IF($B{r}="","",SUM(C{r}:N{r}))'), 'num')
        put(ws, f'P{r}', f(f'IF($B{r}="","",{P[f"Participação sócio {k + 1}"]})'), 'pct')
    put(ws, 'B10', 'Total de retiradas', 'tot_rot')
    for i in range(12):
        c = L(3 + i)
        put(ws, f'{c}10', f(f'IF({fechado(JAN[0] + i)},SUM({c}6:{c}9),"")'), 'tot')
    put(ws, 'O10', f('SUM(O6:O9)'), 'tot')
    put(ws, 'P10', f('SUM(P6:P9)'), 'tot_pct')

    put(ws, 'B12', 'Resultado do ano atribuído', 'secao')
    for k in range(4):
        r = 13 + k
        put(ws, f'B{r}', f(f'B{6 + k}'), 'rot')
        col = L(29 + k)  # AC..AF
        for i in range(12):
            rd = JAN[0] + i
            put(ws, f'{L(3 + i)}{r}', f(f'IF(OR($B{r}="",Dados!$C${rd}<>1),"",{dref(col, rd)})'), 'num')
        put(ws, f'O{r}', f(f'IF($B{r}="","",{vmes(col)})'), 'num')
    put(ws, 'B17', 'Total atribuído', 'tot_rot')
    for i in range(12):
        c = L(3 + i)
        put(ws, f'{c}17', f(f'IF({fechado(JAN[0] + i)},SUM({c}13:{c}16),"")'), 'tot')
    put(ws, 'O17', f('SUM(O13:O16)'), 'tot')
    destacar_mes(ws, 'C6:N17')
    ws.freeze_panes = 'C5'
    nota(ws, 'B18', 'Na coluna Ano, o resultado atribuído é o do mês da competência (já acumula o ano).')
    return ws


def aba_redes(wb):
    ws = nova_aba(wb, 'Redes', f('"Redes sociais · "&' + ANO),
                  'Seguidores no fim do mês, informados no fechamento · "s/d": mês fechado sem o número informado',
                  TERRA, largura_rotulo=26, logo_col='N')
    larguras(ws, mes_cols(), 9)
    ws.column_dimensions['O'].width = 12
    cab_meses(ws, 4, rotulo='Seguidores')
    put(ws, 'O4', 'Ganho no ano', 'cab')

    def seg(col, row):
        return f'IF({fechado(row)},IF(ISNUMBER({dref(col, row)}),{dref(col, row)},"s/d"),"")'

    for r, rede, col in [(5, 'LinkedIn', 'W'), (7, 'Instagram', 'X')]:
        linha_mensal(ws, r, rede, lambda i, rd, c, col=col: seg(col, rd), 'int')
        ws[f'B{r}'].font = fonte(bold=True)
        ant = dref(col, JAN[0] - 1)

        def ganho(i, rd, c, r=r, ant=ant):
            prev = ant if i == 0 else f'{L(2 + i)}{r}'
            return f'IF(AND(ISNUMBER({c}{r}),ISNUMBER({prev})),{c}{r}-{prev},"")'
        linha_mensal(ws, r + 1, 'ganho no mês', ganho, 'int', 'rot_sub')
        for c in mes_cols():
            ws[f'{c}{r + 1}'].font = fonte(color=APAGADO)
        put(ws, f'O{r}', f(f'IF(ISNUMBER({ant}),IFERROR(LOOKUP(2,1/ISNUMBER(C{r}:N{r}),C{r}:N{r})-{ant},""),"")'), 'int')
        put(ws, f'O{r + 1}', None, 'int')
    put(ws, 'B10', f('"Ano anterior · "&(' + ANO + '-1)'), 'secao')
    for r, rede, col in [(11, 'LinkedIn', 'W'), (12, 'Instagram', 'X')]:
        put(ws, f'B{r}', rede, 'rot')
        for i in range(12):
            put(ws, f'{L(3 + i)}{r}', f(seg(col, JAN[-1] + i)), 'int')
        put(ws, f'O{r}', None, 'int')
    ws.conditional_formatting.add('C5:N12', FormulaRule(formula=['C5="s/d"'], font=Font(color=APAGADO, italic=True)))
    destacar_mes(ws, 'C5:N8')
    for c in mes_cols():
        for r in (5, 7, 11, 12):
            ws[f'{c}{r}'].alignment = Alignment(horizontal='right')
    ws.freeze_panes = 'C5'

    ln = LineChart()
    ln.series.append(serie(ws, auxiliar(ws, 5), 3, 14, '$B$5', VINHO, 'line'))
    ln.series.append(serie(ws, auxiliar(ws, 7, 1), 3, 14, '$B$7', TERRA, 'line'))
    rotulos(ln, ws, 3, 4, 14, 4)
    estilo_grafico(ln, 'Seguidores no fim do mês')
    mostrar_ocultas(ln)
    ln.y_axis.scaling.min = 0
    ln.width, ln.height = 24, 8
    ws.add_chart(ln, 'B14')
    return ws


def aba_dados(wb):
    ws = wb.create_sheet('Dados')
    ws.sheet_properties.tabColor = BEGE
    ws.sheet_view.showGridLines = True
    put(ws, 'A1', 'Dados do fechamento', 'titulo')
    put(ws, 'A2', 'Um mês por linha, de dez/(ano-3) a dez/(ano), com centavos; à direita, os parâmetros e as listas usadas nas abas; abaixo, o detalhe do mês por cliente, estado, categoria e fornecedor. Todas as abas leem daqui.', 'sub')
    ws.row_dimensions[D_CAB].height = 30
    for col, nome, fmt in COLS:
        put(ws, f'{col}{D_CAB}', nome, 'dados_cab')
        ws.column_dimensions[col].width = 11 if col not in ('A', 'B', 'C') else 8
        for r in range(D_INI, D_FIM + 1):
            put(ws, f'{col}{r}', None, 'dados', fmt=fmt)
    ws.column_dimensions['C'].width = 7.5
    put(ws, f'AH{D_CAB}', 'Parâmetro', 'dados_cab')
    put(ws, f'AI{D_CAB}', 'Valor', 'dados_cab')
    ws.column_dimensions['AH'].width = 34
    ws.column_dimensions['AI'].width = 16
    for i, (nome, fmt) in enumerate(PARAM):
        put(ws, f'AH{5 + i}', nome, 'dados')
        put(ws, f'AI{5 + i}', None, 'dados', fmt=fmt)
    for col, nome in LISTAS:
        put(ws, f'{col}{D_CAB}', nome, 'dados_cab')
        ws.column_dimensions[col].width = 26 if col in ('AK', 'AO', 'AQ', 'AS') else 9
        for r in range(5, 5 + SLOTS_LISTA):
            put(ws, f'{col}{r}', None, 'dados')
    for col in ('AJ', 'AM', 'AP', 'AR'):
        ws.column_dimensions[col].width = 3
    put(ws, f'A{DET_CAB - 1}', 'Detalhe do mês', 'secao')
    for col, nome, fmt in DET:
        put(ws, f'{col}{DET_CAB}', nome, 'dados_cab')
        put(ws, f'{col}{DET_INI}', None, 'dados', fmt=fmt)
    ws.freeze_panes = 'D5'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    return ws


def main():
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'public/workbook_fechamento_v2.xlsx'
    wb = Workbook()
    normal = wb._named_styles['Normal']
    normal.font = Font(name=FONTE, size=9, color=TINTA)
    aba_resumo(wb)
    aba_resultado(wb)
    aba_faturamento(wb)
    aba_caixa(wb)
    aba_despesas(wb)
    aba_clientes(wb)
    aba_estados(wb)
    aba_capital(wb)
    aba_estoque(wb)
    aba_socios(wb)
    aba_redes(wb)
    aba_dados(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.properties.title = 'Workbook de Fechamento'
    wb.properties.creator = 'AviZee'
    wb.save(saida)
    print(f'ok: {saida}')


if __name__ == '__main__':
    main()
